"""Listes de classe : recueillir les noms des eleves par un QR code.

L'enseignant cree une liste (« B2 CDA ») et en projette le QR code. Chaque
eleve le flashe avec son telephone, saisit son nom et son prenom, et
valide. L'enseignant voit les inscriptions arriver, garde la liste aussi
longtemps qu'il veut, et l'exporte en Excel.

Rien a voir avec une session d'exercices : pas de jeton, pas de cookie,
pas de copie. Une liste n'est qu'un releve de noms horodates.
"""

import io
import random
import re
from datetime import datetime
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from flask import (Blueprint, Response, abort, jsonify, redirect,
                   render_template, request, url_for)
from openpyxl import Workbook
from openpyxl.styles import Font
from werkzeug.utils import secure_filename

from .admin import CODE_ALPHABET, require_admin, who
from .db import execute, now, query

bp = Blueprint("roster", __name__)

# La base garde l'UTC, comme partout ailleurs ; l'enseignant, lui, lit
# l'heure de sa salle, a l'ecran comme dans le classeur.
LOCAL_TZ = "Europe/Paris"

# Le lien circule au-dela de la salle des qu'il est projete : une borne
# evite qu'un plaisantin remplisse la liste sans fin.
MAX_ENTRIES = 500
NAME_MAX = 60
TITLE_MAX = 120


def clean(value, limit=NAME_MAX):
    """Espaces superflus retires, longueur bornee : comme a l'inscription
    d'une session (cf. `student.join`)."""
    return " ".join((value or "").split())[:limit]


def local_time(stamp):
    """Horodatage UTC de la base -> date et heure de la salle, sans fuseau.

    Sans fuseau : openpyxl n'ecrit pas de date qui en porte un, et
    l'enseignant n'en a que faire. Faute de base des fuseaux sur la
    machine, l'heure reste en UTC plutot que de tomber en erreur.
    """
    moment = datetime.fromisoformat(stamp)
    try:
        moment = moment.astimezone(ZoneInfo(LOCAL_TZ))
    except ZoneInfoNotFoundError:
        pass
    return moment.replace(tzinfo=None)


@bp.app_template_filter("local_date")
def local_date(stamp):
    """Filtre de gabarit : « 2026-10-15T09:14:38+00:00 » -> « 15/10/2026 »."""
    return local_time(stamp).strftime("%d/%m/%Y") if stamp else ""


def new_code():
    for _ in range(50):
        code = "".join(random.choice(CODE_ALPHABET) for _ in range(6))
        if query("SELECT 1 FROM roster WHERE code = ?", (code,),
                 one=True) is None:
            return code
    raise RuntimeError("Impossible de generer un code de liste libre.")


def by_code(code):
    return query("SELECT * FROM roster WHERE code = ?",
                 (code.strip().upper()[:12],), one=True)


def get_roster(roster_id):
    roster = query("SELECT * FROM roster WHERE id = ?", (roster_id,), one=True)
    if roster is None:
        abort(404, description="Liste introuvable.")
    return roster


def entries_of(roster_id):
    """Les inscrits, par ordre alphabetique : c'est l'ordre d'une liste
    d'appel, et celui dans lequel l'enseignant cherche un nom."""
    return query(
        """SELECT * FROM roster_entry WHERE roster_id = ?
            ORDER BY last_name COLLATE NOCASE, first_name COLLATE NOCASE,
                     id""",
        (roster_id,),
    )


def listing():
    """Toutes les listes, la plus recente en tete, avec leur effectif."""
    return query(
        """SELECT r.*, COUNT(e.id) AS entries
             FROM roster r LEFT JOIN roster_entry e ON e.roster_id = r.id
            GROUP BY r.id ORDER BY r.created_at DESC, r.id DESC""",
    )


# --------------------------------------------------------------------------
# Cote eleve : la page ouverte par le QR code
# --------------------------------------------------------------------------

@bp.get("/l/<code>")
def form(code):
    roster = by_code(code)
    if roster is None:
        return render_template("roster_join.html", roster=None), 404
    return render_template("roster_join.html", roster=roster)


@bp.post("/l/<code>")
def register(code):
    roster = by_code(code)
    if roster is None:
        return render_template("roster_join.html", roster=None), 404
    if roster["status"] != "open":
        return render_template("roster_join.html", roster=roster)

    first = clean(request.form.get("first_name"))
    last = clean(request.form.get("last_name"))
    if not first or not last:
        return render_template(
            "roster_join.html", roster=roster, first=first, last=last,
            error="Le nom et le prénom sont obligatoires.")

    already = query(
        """SELECT 1 FROM roster_entry
            WHERE roster_id = ? AND last_name = ? COLLATE NOCASE
              AND first_name = ? COLLATE NOCASE""",
        (roster["id"], last, first), one=True)
    if not already:
        count = query("SELECT COUNT(*) AS c FROM roster_entry "
                      "WHERE roster_id = ?", (roster["id"],), one=True)["c"]
        if count >= MAX_ENTRIES:
            return render_template(
                "roster_join.html", roster=roster, first=first, last=last,
                error="Cette liste est complète. Prévenez l'enseignant.")
        execute(
            """INSERT OR IGNORE INTO roster_entry
                      (roster_id, first_name, last_name, created_at)
               VALUES (?, ?, ?, ?)""",
            (roster["id"], first, last, now()))

    # Pas de redirection : la page de confirmation n'a pas d'adresse a
    # elle, et la renvoyer en rechargeant ne cree pas de doublon.
    return render_template("roster_join.html", roster=roster, first=first,
                           last=last, done=True, already=bool(already))


# --------------------------------------------------------------------------
# Cote enseignant
# --------------------------------------------------------------------------

@bp.post("/admin/listes")
@require_admin
def create():
    title = clean(request.form.get("title"), TITLE_MAX)
    if not title:
        return redirect(url_for("admin.dashboard",
                                error="Donnez un nom à la liste."))
    roster_id = execute(
        """INSERT INTO roster (code, title, status, created_by, created_at)
           VALUES (?, ?, 'open', ?, ?)""",
        (new_code(), title, who(), now()))
    return redirect(url_for("roster.view", roster_id=roster_id))


@bp.get("/admin/listes/<int:roster_id>")
@require_admin
def view(roster_id):
    roster = get_roster(roster_id)
    return render_template(
        "admin_roster.html", roster=roster,
        join_url=url_for("roster.form", code=roster["code"], _external=True))


@bp.get("/admin/api/listes/<int:roster_id>")
@require_admin
def api_entries(roster_id):
    """Les inscrits, relus par la page de l'enseignant toutes les 3 s."""
    roster = get_roster(roster_id)
    rows = entries_of(roster_id)
    return jsonify({
        "status": roster["status"],
        "count": len(rows),
        "entries": [{
            "id": row["id"],
            "last_name": row["last_name"],
            "first_name": row["first_name"],
            "at": local_time(row["created_at"]).strftime("%d/%m/%Y %H:%M:%S"),
        } for row in rows],
    })


@bp.post("/admin/listes/<int:roster_id>/close")
@require_admin
def close(roster_id):
    get_roster(roster_id)
    execute("""UPDATE roster SET status = 'closed', closed_at = ?
                WHERE id = ? AND status = 'open'""", (now(), roster_id))
    return redirect(url_for("roster.view", roster_id=roster_id))


@bp.post("/admin/listes/<int:roster_id>/open")
@require_admin
def reopen(roster_id):
    get_roster(roster_id)
    execute("""UPDATE roster SET status = 'open', closed_at = NULL
                WHERE id = ?""", (roster_id,))
    return redirect(url_for("roster.view", roster_id=roster_id))


@bp.post("/admin/listes/<int:roster_id>/delete")
@require_admin
def delete(roster_id):
    roster = get_roster(roster_id)
    execute("DELETE FROM roster WHERE id = ?", (roster_id,))
    return redirect(url_for("admin.dashboard",
                            ok="Liste « %s » supprimée." % roster["title"]))


@bp.post("/admin/listes/<int:roster_id>/entrees/<int:entry_id>/delete")
@require_admin
def delete_entry(roster_id, entry_id):
    """Retire une ligne : un doublon mal orthographie, un nom farfelu."""
    get_roster(roster_id)
    execute("DELETE FROM roster_entry WHERE id = ? AND roster_id = ?",
            (entry_id, roster_id))
    return jsonify({"ok": True})


def workbook(roster, rows):
    """Le classeur de la liste : nom, prenom, date et heure d'inscription."""
    book = Workbook()
    sheet = book.active
    # Un nom de feuille ne tolere ni `[]:*?/\` ni plus de 31 caracteres.
    sheet.title = re.sub(r"[\[\]:*?/\\]", " ", roster["title"])[:31].strip() \
        or "Liste"
    sheet.append(["Nom", "Prénom", "Date et heure d'enregistrement"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([row["last_name"], row["first_name"],
                      local_time(row["created_at"])])
        line = sheet.max_row
        # Un nom saisi par l'eleve reste du texte : sans cela, « =1+1 »
        # deviendrait une formule a l'ouverture du classeur.
        for col in (1, 2):
            sheet.cell(line, col).data_type = "s"
        sheet.cell(line, 3).number_format = "dd/mm/yyyy hh:mm:ss"
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 20
    sheet.column_dimensions["C"].width = 30
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:C%d" % max(sheet.max_row, 1)
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


@bp.get("/admin/listes/<int:roster_id>/export.xlsx")
@require_admin
def export_xlsx(roster_id):
    roster = get_roster(roster_id)
    name = secure_filename("liste-%s" % roster["title"]) or "liste"
    return Response(
        workbook(roster, entries_of(roster_id)),
        mimetype="application/vnd.openxmlformats-officedocument."
                 "spreadsheetml.sheet",
        headers={"Content-Disposition":
                 'attachment; filename="%s.xlsx"' % name},
    )
